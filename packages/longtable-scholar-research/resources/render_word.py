import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def set_run_font(run, latin, east_asia, size):
    run.font.name = latin
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instruction, separate, text, end):
        run._r.append(item)


def style_document(doc, profile):
    section = doc.sections[0]
    section.page_width = Inches(8.5 if profile["page"]["size"] == "Letter" else 8.27)
    section.page_height = Inches(11 if profile["page"]["size"] == "Letter" else 11.69)
    margin = Inches(profile["page"]["marginInches"])
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = margin
    section.header_distance = Inches(profile["page"]["headerInches"])
    section.footer_distance = Inches(profile["page"]["footerInches"])
    body = profile["body"]
    normal = doc.styles["Normal"]
    normal.font.name = body["latinFont"]
    normal.font.size = Pt(body["fontSizePt"])
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), body["latinFont"])
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), body["latinFont"])
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), body["eastAsiaFont"])
    normal.paragraph_format.line_spacing = body["lineSpacing"]
    normal.paragraph_format.space_after = Pt(body["paragraphAfterPt"])
    for style_name in ("Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = body["latinFont"]
        style.font.size = Pt(body["fontSizePt"])
        style.font.bold = True
        style.font.color.rgb = None
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), body["eastAsiaFont"])
        style.paragraph_format.line_spacing = body["lineSpacing"]
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
    page_paragraph = section.header.paragraphs[0] if profile["pageNumber"] == "top_right" else section.footer.paragraphs[0]
    page_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if profile["pageNumber"] == "top_right" else WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(page_paragraph)


def add_centered(doc, text, bold=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    return paragraph


def add_body(doc, text, profile, indent=True):
    paragraph = doc.add_paragraph()
    if indent:
        paragraph.paragraph_format.first_line_indent = Inches(profile["body"]["firstLineIndentInches"])
    paragraph.add_run(text)
    return paragraph


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def add_readiness_table(doc, manuscript):
    rows = [
        ("점검 항목", "상태"),
        ("포함 연구", f'{len(manuscript["references"])}편'),
        ("효과크기 산출 가능성", "승인된 분석계획에 따라 별도 추출"),
        ("표본 및 출판물 의존성", "코딩 필요"),
        ("결측 및 저자 연락", "코딩 필요"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value
            if row_index == 0:
                for run in table.cell(row_index, column_index).paragraphs[0].runs:
                    run.bold = True
    set_table_geometry(table, [3000, 6360])


def add_markdown_italic_runs(paragraph, text):
    for index, part in enumerate(re.split(r"\*([^*]+)\*", text)):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.italic = index % 2 == 1


def build(payload, output_path):
    manuscript = payload["manuscript"]
    profile = payload["profile"]
    template_path = profile.get("templatePath")
    doc = Document(template_path) if template_path else Document()
    style_document(doc, profile)
    if profile["titlePage"]:
        for _ in range(5):
            doc.add_paragraph()
        add_centered(doc, manuscript["title"], bold=True)
        doc.add_paragraph()
        for line in manuscript["authorLines"]:
            add_centered(doc, line)
        doc.add_page_break()
    else:
        add_centered(doc, manuscript["title"], bold=True)
        for line in manuscript["authorLines"]:
            add_centered(doc, line)
    add_centered(doc, "초록" if manuscript["language"] == "ko" else "Abstract", bold=True)
    add_body(doc, manuscript["abstract"], profile, indent=False)
    keywords = doc.add_paragraph()
    label = keywords.add_run("주제어: " if manuscript["language"] == "ko" else "Keywords: ")
    label.italic = True
    keywords.add_run(", ".join(manuscript["keywords"]))
    keywords.paragraph_format.space_after = Pt(12)
    for section in manuscript["sections"]:
        heading = doc.add_paragraph(style="Heading 1")
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(0)
        heading.paragraph_format.keep_with_next = True
        heading.add_run(section["heading"])
        add_body(doc, section["content"], profile)
        if section["id"] == "methods":
            add_readiness_table(doc, manuscript)
    add_centered(doc, "참고문헌" if manuscript["language"] == "ko" else "References", bold=True)
    hanging = profile["references"]["hangingIndentInches"]
    for reference in manuscript["references"]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(hanging)
        paragraph.paragraph_format.first_line_indent = Inches(-hanging)
        add_markdown_italic_runs(paragraph, reference)
    doc.core_properties.title = manuscript["title"]
    doc.core_properties.subject = f'LongTable protocol {manuscript["protocol"]["id"]}'
    doc.core_properties.comments = f'Render profile {profile["id"]}; generated {manuscript["generatedAt"]}'
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


if __name__ == "__main__":
    build(json.load(sys.stdin), sys.argv[1])
